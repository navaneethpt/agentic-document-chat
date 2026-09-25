from io import BytesIO

from docx import Document
from pypdf import PdfWriter
import pytest
from reportlab.pdfgen.canvas import Canvas

from rag_chat.documents import DocumentError, Section, chunk_sections, extract_sections, MAX_FILE_BYTES


@pytest.mark.parametrize("extension", ["txt", "md"])
def test_text(extension):
    sections = extract_sections(f"notes.{extension}", b"\xef\xbb\xbfLaunch is in June.")
    assert sections == [Section("Launch is in June.", "Text")]


def test_pdf_pages():
    buffer = BytesIO()
    canvas = Canvas(buffer)
    for text in ["First page", "Second page"]:
        canvas.drawString(60, 700, text)
        canvas.showPage()
    canvas.save()
    sections = extract_sections("pages.pdf", buffer.getvalue())
    assert [section.location for section in sections] == ["Page 1", "Page 2"]
    assert "Second page" in sections[1].text


def test_docx_preserves_table_order():
    document = Document()
    document.add_paragraph("Before")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Launch"
    table.cell(0, 1).text = "June"
    document.add_paragraph("After")
    buffer = BytesIO()
    document.save(buffer)
    sections = extract_sections("notes.docx", buffer.getvalue())
    assert [section.text for section in sections] == ["Before", "Launch | June", "After"]
    assert sections[1].location == "Table/block 2, row 1"


@pytest.mark.parametrize("name,data,match", [
    ("x.exe", b"hello", "Unsupported"),
    ("x.txt", b"", "empty"),
    ("x.txt", b"\xff", "UTF-8"),
    ("x.txt", b" \n\t", "No extractable"),
    ("x.pdf", b"broken", "Cannot read"),
    ("x.docx", b"broken", "Cannot read"),
    ("x.txt", b"x" * (MAX_FILE_BYTES + 1), "20 MB"),
])
def test_bad_uploads(name, data, match):
    with pytest.raises(DocumentError, match=match):
        extract_sections(name, data)


def test_scanned_and_encrypted_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(DocumentError, match="OCR"):
        extract_sections("scan.pdf", buffer.getvalue())
    writer.encrypt("secret")
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(DocumentError, match="Password-protected"):
        extract_sections("locked.pdf", buffer.getvalue())


def test_long_text_chunk_boundaries(tokenizer):
    words = [f"word{i}" for i in range(530)]
    chunks = chunk_sections([Section(" ".join(words), "Page 3")], "../notes.pdf", tokenizer)
    assert [len(chunk.text.split()) for chunk in chunks] == [200, 200, 200, 50]
    assert chunks[0].text.split()[-40:] == chunks[1].text.split()[:40]
    assert chunks[-1].text.endswith("word529")
    assert all(chunk.filename == "notes.pdf" and chunk.location == "Page 3" for chunk in chunks)


def test_invalid_chunk_settings(tokenizer):
    with pytest.raises(ValueError):
        chunk_sections([Section("text", "Text")], "x.txt", tokenizer, size=10, overlap=10)
